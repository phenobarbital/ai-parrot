"""BBVA movements export → :class:`BbvaStatement` (FEAT-602 M8)."""

from __future__ import annotations

import asyncio
import datetime as dt
import logging
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import openpyxl
from parrot_loaders.excel import ExcelLoader
from parrot_tools.business_automation.ingest import compute_statement_digest

from ..models import BankExpenseRow, BbvaStatement

logger = logging.getLogger(__name__)

HEADER_TOKENS = {
    "fecha",
    "f.valor",
    "fecha valor",
    "concepto",
    "movimiento",
    "importe",
    "divisa",
    "disponible",
    "observaciones",
}

#: Minimum number of matching tokens for a scanned row to be considered the header row.
_MIN_HEADER_MATCHES = 3
#: How many rows (from the top of the sheet) are scanned looking for the header.
_HEADER_SCAN_LIMIT = 40


def _normalize_token(value: Any) -> str:
    """Case- and accent-insensitive normalization of a header cell value."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text.strip().lower()


def parse_es_amount(value: Any) -> Optional[Decimal]:
    """``-1.234,56`` / ``1.234,56 €`` / numeric → Decimal; blank → None."""
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        # bool is an int subclass; a BBVA amount cell is never a boolean.
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))

    text = str(value).strip()
    if not text:
        return None

    text = text.replace("€", "").strip()
    negative = False
    if text.startswith("-"):
        negative = True
        text = text[1:].strip()
    elif text.startswith("+"):
        text = text[1:].strip()

    # Spanish locale: "." thousands separator, "," decimal separator.
    text = text.replace(".", "").replace(",", ".")
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Cannot parse Spanish amount: {value!r}") from exc
    return -amount if negative else amount


def parse_es_date(value: Any) -> Optional[dt.date]:
    """datetime/date cell or ``dd/mm/yyyy`` → date; blank → None."""
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value

    text = str(value).strip()
    if not text:
        return None

    try:
        return dt.datetime.strptime(text, "%d/%m/%Y").date()
    except ValueError as exc:
        raise ValueError(f"Cannot parse Spanish date: {value!r}") from exc


def _parse_bbva_sync(path: Path, sheet: Optional[str]) -> BbvaStatement:
    """Pure synchronous parse (runs in a worker thread)."""
    digest = compute_statement_digest(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
        sheet_name = worksheet.title

        max_scan = min(_HEADER_SCAN_LIMIT, worksheet.max_row or 0)
        header_row_index: Optional[int] = None
        header_values: tuple = ()
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=1, max_row=max_scan, values_only=True)):
            tokens = {_normalize_token(cell) for cell in row}
            matches = tokens & HEADER_TOKENS
            if len(matches) >= _MIN_HEADER_MATCHES and "fecha" in tokens and "importe" in tokens:
                header_row_index = row_idx
                header_values = row
                break

        if header_row_index is None:
            raise ValueError(
                f"BBVA header row not found in the first {max_scan} rows of sheet {sheet_name!r} in {path}"
            )

        column_map: Dict[str, int] = {}
        for col_idx, cell in enumerate(header_values):
            token = _normalize_token(cell)
            if token:
                column_map[token] = col_idx

        def _col(*keys: str) -> Optional[int]:
            for key in keys:
                if key in column_map:
                    return column_map[key]
            return None

        booking_col = _col("fecha")
        value_col = _col("f.valor", "fecha valor")
        concept_col = _col("concepto")
        movement_col = _col("movimiento")
        amount_col = _col("importe")
        currency_col = _col("divisa")
        balance_col = _col("disponible")
        observations_col = _col("observaciones")

        def _text(row: tuple, col: Optional[int]) -> Optional[str]:
            if col is None or row[col] is None:
                return None
            text = str(row[col]).strip()
            return text or None

        rows: List[BankExpenseRow] = []
        skipped = 0
        row_count = 0

        data_start = header_row_index + 2  # 1-based Excel row right after the header row.
        for offset, row in enumerate(
            worksheet.iter_rows(min_row=data_start, max_row=worksheet.max_row, values_only=True)
        ):
            row_count += 1

            amount = parse_es_amount(row[amount_col]) if amount_col is not None else None
            booking_date = parse_es_date(row[booking_col]) if booking_col is not None else None

            if amount is None or amount >= 0 or booking_date is None:
                skipped += 1
                continue

            rows.append(
                BankExpenseRow(
                    row_index=offset,
                    booking_date=booking_date,
                    value_date=parse_es_date(row[value_col]) if value_col is not None else None,
                    concept=_text(row, concept_col) or "",
                    movement=_text(row, movement_col),
                    amount=amount,
                    currency=_text(row, currency_col) or "EUR",
                    balance=parse_es_amount(row[balance_col]) if balance_col is not None else None,
                    observations=_text(row, observations_col),
                    row_id=f"{digest}:{offset}",
                )
            )

        return BbvaStatement(
            path=str(path),
            digest=digest,
            sheet=sheet_name,
            header_row=header_row_index,
            rows=rows,
            skipped=skipped,
            row_count=row_count,
        )
    finally:
        workbook.close()


async def parse_bbva_statement(path: Union[str, Path], *, sheet: Optional[str] = None) -> BbvaStatement:
    """Parse a BBVA export off the event loop and cross-check its row count with ExcelLoader."""
    path = Path(path)
    statement = await asyncio.to_thread(_parse_bbva_sync, path, sheet)

    # Independent row-count cross-check (S10): ExcelLoader's row mode is
    # told the same detected header row via its ``header=`` constructor
    # kwarg, which pandas honours by skipping rows 0..header-1 and using
    # ``header`` as the column row — the same semantics the preamble
    # requires, so no separate raw-pandas fallback is needed here (see the
    # Completion Note for the rationale).
    loader = ExcelLoader(path, output_mode="row", header=statement.header_row, sheets=statement.sheet)
    documents = await loader.load(path, split_documents=False)

    if len(documents) != statement.row_count:
        raise ValueError(
            f"Row count mismatch between BBVA parser ({statement.row_count}) and "
            f"ExcelLoader ({len(documents)}) for {path}"
        )
    return statement
